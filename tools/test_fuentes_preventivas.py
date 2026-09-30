import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from datetime import datetime
import openpyxl

MODULE = Path(__file__).with_name('fuentes_preventivas.py')
spec = importlib.util.spec_from_file_location('fuentes', MODULE)
fuentes = importlib.util.module_from_spec(spec) if MODULE.exists() else None
if fuentes:
    spec.loader.exec_module(fuentes)


class FuentesTest(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(fuentes, 'Falta el adaptador de fuentes reales')
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def libro(self, nombre, headers, rows, extra=None):
        p = Path(self.tmp.name) / 'fuente.xlsx'
        w = openpyxl.Workbook(); s = w.active; s.title = nombre
        s.append(headers)
        for row in rows: s.append(row)
        if extra:
            x = w.create_sheet('Resumen'); x.append(extra)
        w.save(p); w.close()
        return p

    def nominal(self, rows):
        return self.libro('DETECCIÓN', ['clues','curpprestador','nombreprestador','fechadeteccion','alcohol','tabaco','cannabis','cocaina','metanfetaminas','inhalables','opiaceos','alucinogenos','tranquilizantes','otras_substancias','registro_eliminado','curppaciente','nombre'], rows)

    def test_cuenta_negativos_una_aplicacion_drogas_y_excluye_eliminados(self):
        row = ['U1','PROVEEDOR_FICTICIO','Recurso ficticio',datetime(2026,9,10),'NEGATIVO','POSITIVO','NEGATIVO','POSITIVO']+['N/A']*6+['NO','PACIENTE_SECRETO','Nombre secreto']
        p = self.nominal([row, row[:3]+[datetime(2026,8,10)]+row[4:],row[:14]+['SI']+row[15:]])
        r = fuentes.nominal(p,'2026-09-28','JS19',b'clave-ficticia-suficientemente-larga')
        sept = [f for f in r['filas'] if f['periodo']=='2026-09']
        self.assertEqual({f['instrumento']:f['total'] for f in sept},{'AUDIT':1,'FAGERSTROM':1,'DROGAS':1})
        self.assertEqual({f['periodo'] for f in r['filas']},{'2026-08','2026-09'})
        salida = json.dumps(r)
        for secreto in ('PACIENTE_SECRETO','Nombre secreto','PROVEEDOR_FICTICIO','curppaciente'):
            self.assertNotIn(secreto,salida)
        self.assertEqual(r['auditoria']['eliminados'],1)

    def test_mismo_paciente_en_dos_registros_no_se_deduplica_por_curp(self):
        row=['U1','R1','Recurso',datetime(2026,9,10),'NEGATIVO','N/A']+['N/A']*8+['NO','PAC1','Paciente']
        r=fuentes.nominal(self.nominal([row,row]),'2026-09-28','JS19',b'clave-ficticia-suficientemente-larga')
        self.assertEqual(next(f['total'] for f in r['filas'] if f['instrumento']=='AUDIT'),2)
        self.assertEqual(r['auditoria']['filas_identicas'],1)

    def test_recurso_ausente_conserva_cantidad_por_unidad_y_marca_pendiente(self):
        row=['U1',None,None,datetime(2026,9,10),'NEGATIVO','N/A']+['N/A']*8+['NO','PAC1','Paciente']
        r=fuentes.nominal(self.nominal([row]),'2026-09-28','JS19',b'clave-ficticia-suficientemente-larga')
        self.assertEqual(next(f['total'] for f in r['filas'] if f['instrumento']=='AUDIT'),1)
        self.assertEqual(r['filas'][0]['estado'],'RECURSO_PENDIENTE')
        self.assertEqual(r['auditoria']['filas_sin_recurso'],1)

    def test_codigo_desconocido_detiene_sin_exponer_la_fila(self):
        row=['U1','R1','Recurso',datetime(2026,9,10),'NO SABEMOS','N/A']+['N/A']*8+['NO','PAC1','Paciente']
        with self.assertRaisesRegex(ValueError,'resultado'):
            fuentes.nominal(self.nominal([row]),'2026-09-28','JS19',b'clave-ficticia-suficientemente-larga')
        for value in (0,False,1):
            with self.assertRaisesRegex(ValueError,'resultado'): fuentes.resultado(value)

    def test_cubos_suma_sexo_resultado_no_resumen_y_no_inventa_recurso(self):
        rows=[['U1',c,1,None,0] for c in ('DET11','DET21','DET35','DET44')]
        p=self.libro('Datos',['CLUES','claveVariable','Enero','Febrero','Septiembre'],rows,extra=['Total',9000])
        r=fuentes.cubos(p,'2026-09-28','JS19',2026)
        f=next(f for f in r['filas'] if f['periodo']=='2026-01')
        self.assertEqual(f['total'],4); self.assertEqual(f['instrumento'],'AUDIT')
        self.assertNotIn('recurso_id',f)
        self.assertNotIn('2026-09',{f['periodo'] for f in r['filas']})
        self.assertNotIn('DROGAS',{f['instrumento'] for f in r['filas']})

    def test_cubos_rechaza_variable_duplicada_y_conserva_categoria_faltante(self):
        p=self.libro('Datos',['CLUES','claveVariable','Enero'],[['U1','DET11',2]])
        r=fuentes.cubos(p,'2026-02-01','JS19',2026)
        self.assertEqual(r['filas'][0]['estado'],'CATEGORIAS_INCOMPLETAS')
        p=self.libro('Datos',['CLUES','claveVariable','Enero'],[['U1','DET11',2],['U1','DET11',2]])
        with self.assertRaisesRegex(ValueError,'duplicada'): fuentes.cubos(p,'2026-02-01','JS19',2026)

    def test_sinba_excluye_subtotales_y_preserva_rango_acumulado(self):
        p=Path(self.tmp.name)/'sis.xlsx';w=openpyxl.Workbook();s=w.active
        s.append(['Fecha inicial','01/01/2026']);s.append(['Fecha final','31/08/2026']);s.append(['Agrupar por','Profesional']);s.append([])
        s.append(['Unidad Médica','Profesional','Detecciones'])
        s.append(['MCSSA000001 Unidad ficticia',None,999]);s.append(['MCSSA000001 Unidad ficticia','AAAA000101HMCBBB01 Recurso ficticio',4])
        w.save(p);w.close()
        r=fuentes.sinba(p,'JS19',b'clave-ficticia-suficientemente-larga')
        self.assertEqual(len(r['filas']),1);self.assertEqual(r['filas'][0]['total'],4)
        self.assertEqual(r['filas'][0]['periodo'],'2026-01/2026-08')
        self.assertEqual(r['filas'][0]['medicion'],'DETECCIONES_GENERALES')
        self.assertNotIn('instrumento',r['filas'][0])

    def test_sinba_recurso_repetido_no_suma_ni_certifica_un_total(self):
        p=Path(self.tmp.name)/'sis.xlsx';w=openpyxl.Workbook();s=w.active
        for r in [['Fecha inicial','01/01/2026'],['Fecha final','31/08/2026'],['Agrupar por','Profesional'],[],['Unidad Médica','Profesional','Detecciones'],['MCSSA000001 Unidad','AAAA000101HMCBBB01 Recurso',4],['MCSSA000001 Unidad','AAAA000101HMCBBB01 Recurso',4]]:s.append(r)
        w.save(p);w.close();r=fuentes.sinba(p,'JS19',b'clave-ficticia-suficientemente-larga')
        self.assertEqual(len(r['filas']),1);self.assertIsNone(r['filas'][0]['total'])
        self.assertEqual(r['filas'][0]['estado'],'REVISAR_RECURSO_DUPLICADO')

    def test_sinba_no_exporta_curp_en_nombre_con_minusculas(self):
        p=Path(self.tmp.name)/'sis.xlsx';w=openpyxl.Workbook();s=w.active
        for r in [['Fecha inicial','01/01/2026'],['Fecha final','31/08/2026'],['Agrupar por','Profesional'],[],['Unidad Médica','Profesional','Detecciones'],['MCSSA000001 Unidad','aaaa000101hmcbbb01 Recurso',4]]:s.append(r)
        w.save(p);w.close();r=fuentes.sinba(p,'JS19',b'clave-ficticia-suficientemente-larga')
        self.assertNotIn('aaaa000101hmcbbb01',json.dumps(r).lower())

    def test_sinba_acepta_fila_con_columnas_finales_vacias(self):
        p=Path(self.tmp.name)/'sis.xlsx';w=openpyxl.Workbook();s=w.active
        for r in [['Fecha inicial','01/01/2026'],['Fecha final','31/08/2026'],['Agrupar por','Profesional'],[],['Unidad Médica','Profesional','Detecciones','Otra columna'],['MCSSA000001 Unidad','AAAA000101HMCBBB01 Recurso',4]]:s.append(r)
        w.save(p);w.close();r=fuentes.sinba(p,'JS19',b'clave-ficticia-suficientemente-larga')
        self.assertEqual(len(r['filas']),1);self.assertEqual(r['filas'][0]['total'],4)


if __name__ == '__main__': unittest.main()
