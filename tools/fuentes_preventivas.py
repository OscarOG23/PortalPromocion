"""Lee archivos institucionales y produce cortes agregados para el backend privado.

No modifica los originales. Nunca exporta pacientes ni CURP de prestadores.
La equivalencia de instrumentos fue confirmada por el usuario el 30/09/2026.
"""
import argparse
import hashlib
import hmac
import json
import re
import secrets
from collections import Counter
from datetime import date, datetime
from pathlib import Path
import openpyxl

SUSTANCIAS = ('cannabis','cocaina','metanfetaminas','inhalables','opiaceos','alucinogenos','tranquilizantes','otras_substancias')
VARIABLES = {'AUDIT': ('DET11','DET21','DET35','DET44'), 'FAGERSTROM': ('DET12','DET22','DET36','DET45')}
MESES = ('Enero','Febrero','Marzo','Abril','Mayo','Junio','Julio','Agosto','Septiembre','Octubre','Noviembre','Diciembre')


def texto(value):
    return '' if value is None else str(value).strip()


def fecha(value):
    if isinstance(value, (date, datetime)):
        return value.date() if isinstance(value, datetime) else value
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try: return datetime.strptime(texto(value), fmt).date()
        except ValueError: pass
    raise ValueError('Fecha inválida en la fuente; revisar sin publicar la fila nominal.')


def base(path, fuente, corte, ambito):
    if not re.fullmatch(r'[A-Za-z0-9_-]{1,60}', ambito): raise ValueError('Ámbito inválido.')
    with open(path,'rb') as archivo:
        digest=hashlib.file_digest(archivo,'sha256').hexdigest()
    return {'version':1, 'fuente':fuente, 'ambito':ambito, 'corte':fecha(corte).isoformat(),
            'sha256':digest, 'filas':[], 'auditoria':{}}


def recurso(identity, clave):
    if len(clave)<32: raise ValueError('La clave privada de recursos debe tener al menos 32 bytes.')
    return 'R_' + hmac.new(clave,texto(identity).upper().encode('utf-8'),hashlib.sha256).hexdigest()


def resultado(value):
    v=texto(value).upper()
    if v in ('POSITIVO','NEGATIVO'): return True
    if v in ('','N/A'): return False
    raise ValueError('Código de resultado desconocido; revisar el diccionario de la fuente.')


def cantidad(value):
    if isinstance(value,bool) or not isinstance(value,(int,float)) or value<0 or int(value)!=value:
        raise ValueError('Cifra inválida en la fuente.')
    return int(value)


def nominal(path, corte, ambito, clave):
    out=base(path,'NOMINAL',corte,ambito); grupos={}; vistos=set(); audit=Counter()
    w=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        if 'DETECCIÓN' not in w: raise ValueError('Falta la hoja DETECCIÓN.')
        rows=w['DETECCIÓN'].iter_rows(values_only=True); headers=list(next(rows)); idx={h:i for i,h in enumerate(headers)}
        required=('clues','curpprestador','fechadeteccion','alcohol','tabaco','registro_eliminado')+SUSTANCIAS
        if any(k not in idx for k in required): raise ValueError('Encabezados nominales incompatibles.')
        for row in rows:
            if not any(v is not None for v in row): continue
            audit['filas_leidas']+=1
            get=lambda k: row[idx[k]] if k in idx and idx[k]<len(row) else None
            eliminado=texto(get('registro_eliminado')).upper()
            if eliminado in ('SI','SÍ'): audit['eliminados']+=1; continue
            if eliminado!='NO': raise ValueError('Estado de eliminación desconocido.')
            d=fecha(get('fechadeteccion'))
            if d>fecha(corte): raise ValueError('El corte es anterior a una fecha de atención.')
            unidad=texto(get('clues')).upper(); prestador=texto(get('curpprestador'))
            if not unidad: raise ValueError('Falta CLUES en DETECCIÓN.')
            if not prestador: audit['filas_sin_recurso']+=1
            # Sólo detecta duplicados para revisión. No deduplica aplicaciones por paciente.
            firma=hashlib.sha256(json.dumps(row,default=str,ensure_ascii=False).encode()).digest()
            if firma in vistos: audit['filas_identicas']+=1
            vistos.add(firma)
            rid=recurso(prestador or 'SIN_IDENTIDAD|'+unidad,clave); periodo=d.strftime('%Y-%m')
            nombre=' '.join(texto(get(k)) for k in ('nombreprestador','primerapellidoprestador','segundoapellidoprestador')).strip()
            if not prestador: nombre='SIN IDENTIDAD DEL RECURSO EN LA FUENTE'
            aplicadas={'AUDIT':resultado(get('alcohol')), 'FAGERSTROM':resultado(get('tabaco'))}
            resultados=[resultado(get(k)) for k in SUSTANCIAS]
            aplicadas['DROGAS']=any(resultados)  # Una aplicación, no una por sustancia.
            for instrumento,aplicada in aplicadas.items():
                key=(periodo,unidad,rid,instrumento)
                if key not in grupos: grupos[key]={'periodo':periodo,'unidad_id':unidad,'recurso_id':rid,'recurso_nombre':nombre,'instrumento':instrumento,'total':0,'estado':'OBSERVADO' if prestador else 'RECURSO_PENDIENTE'}
                grupos[key]['total']+=int(aplicada)
        out['filas']=[grupos[k] for k in sorted(grupos)]
        out['auditoria']=dict(audit); out['auditoria'].setdefault('filas_identicas',0)
        return out
    finally:
        if 'rows' in locals(): rows.close()
        w.close()


def cubos(path, corte, ambito, anio):
    if not 2000<=anio<=2099: raise ValueError('Año inválido.')
    out=base(path,'CUBOS',corte,ambito); activos=set(); seleccion=[]; unicos=set()
    w=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        if 'Datos' not in w: raise ValueError('Falta la hoja Datos de Cubos; Resumen anual no es mensual.')
        rows=w['Datos'].iter_rows(values_only=True); headers=list(next(rows));idx={h:i for i,h in enumerate(headers)}
        if any(k not in idx for k in ('CLUES','claveVariable')) or not any(m in idx for m in MESES): raise ValueError('Encabezados de Cubos incompatibles.')
        codes={c:i for i,cs in VARIABLES.items() for c in cs}
        for row in rows:
            if not any(v is not None for v in row): continue
            vals={m:row[idx[m]] for m in MESES if m in idx and idx[m]<len(row)}
            # El exportador rellena meses no disponibles con cero: no se certifican.
            for m,v in vals.items():
                if v not in (None,'',0): cantidad(v); activos.add(m)
            code=texto(row[idx['claveVariable']]);unidad=texto(row[idx['CLUES']]).upper()
            if code not in codes: continue
            if not unidad: raise ValueError('Falta CLUES en Cubos.')
            if (unidad,code) in unicos: raise ValueError('Variable duplicada por CLUES en Cubos; no sumar orígenes/subtotales.')
            unicos.add((unidad,code)); seleccion.append((unidad,code,vals))
        grupos={}
        for unidad,code,vals in seleccion:
            for m in activos:
                v=vals.get(m)
                if v is None or v=='': continue
                periodo=f'{anio}-{MESES.index(m)+1:02d}'
                if periodo>fecha(corte).strftime('%Y-%m'): raise ValueError('Cubos contiene meses posteriores al corte declarado.')
                k=(periodo,unidad,codes[code]);g=grupos.setdefault(k,{'periodo':periodo,'unidad_id':unidad,'instrumento':codes[code],'total':0,'categorias':[]})
                g['total']+=cantidad(v);g['categorias'].append(code)
        for k in sorted(grupos):
            g=grupos[k];g['estado']='OBSERVADO' if len(g['categorias'])==4 else 'CATEGORIAS_INCOMPLETAS';out['filas'].append(g)
        out['auditoria']={'meses_con_actividad':[m for m in MESES if m in activos], 'drogas':'SIN_VARIABLE_EQUIVALENTE'}
        return out
    finally:
        if 'rows' in locals(): rows.close()
        w.close()


def sinba(path, ambito, clave):
    w=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        s=w.active;s.reset_dimensions();it=s.iter_rows(values_only=True); inicio=fecha(next(it)[1]);fin=fecha(next(it)[1]);next(it);next(it)
        headers=list(next(it));idx={h:i for i,h in enumerate(headers)}
        if any(k not in idx for k in ('Unidad Médica','Profesional','Detecciones')): raise ValueError('Encabezados SINBA incompatibles.')
        if fin<inicio: raise ValueError('Intervalo SINBA inválido.')
        out=base(path,'SINBA',fin.isoformat(),ambito);out['inicio']=inicio.isoformat();out['fin']=fin.isoformat()
        periodo=inicio.strftime('%Y-%m')+'/'+fin.strftime('%Y-%m');seen={};audit=Counter()
        for row in it:
            if len(row)<=max(idx[k] for k in ('Unidad Médica','Profesional','Detecciones')): continue
            profesional=texto(row[idx['Profesional']]);unidad=texto(row[idx['Unidad Médica']])
            if not profesional: audit['subtotales_excluidos']+=1;continue
            u=re.search(r'\b[A-Z]{5}\d{6}\b',unidad.upper())
            if not u: raise ValueError('Fila SINBA por recurso sin CLUES; revisar el formato.')
            p=re.search(r'\b[A-Z][AEIOUX][A-Z]{2}\d{6}[HM][A-Z]{5}[A-Z0-9]\d\b',profesional.upper())
            identidad=p.group() if p else 'SIN_CURP|'+profesional.upper()
            if not p: audit['recursos_sin_curp']+=1
            rid=recurso(identidad,clave); k=(u.group(),rid)
            if k in seen:
                seen[k]['total']=None;seen[k]['estado']='REVISAR_RECURSO_DUPLICADO'
                audit['filas_recurso_repetido']+=1
                continue
            v=row[idx['Detecciones']]
            # Vacío no equivale a cero.
            total=None if v in (None,'') else cantidad(v)
            nombre=re.sub(re.escape(p.group()),'',profesional,flags=re.IGNORECASE).strip(' -') if p else profesional
            dato={'periodo':periodo,'unidad_id':u.group(),'recurso_id':rid,'recurso_nombre':nombre,'medicion':'DETECCIONES_GENERALES','total':total,'estado':'OBSERVADO' if p else 'RECURSO_PENDIENTE'}
            seen[k]=dato;out['filas'].append(dato)
        out['auditoria']=dict(audit);return out
    finally:
        if 'it' in locals(): it.close()
        w.close()


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('fuente',choices=('nominal','cubos','sinba'));parser.add_argument('archivo',type=Path)
    parser.add_argument('--ambito',required=True);parser.add_argument('--corte');parser.add_argument('--anio',type=int)
    parser.add_argument('--clave-recursos',type=Path);parser.add_argument('--salida',required=True,type=Path)
    args=parser.parse_args()
    if args.fuente!='sinba' and not args.corte: parser.error('Nominal/Cubos requieren --corte de la fuente, no la fecha de importación.')
    clave=None
    if args.fuente!='cubos':
        if not args.clave_recursos: parser.error('Se requiere --clave-recursos (mismo archivo privado en todos los meses).')
        args.clave_recursos.parent.mkdir(parents=True,exist_ok=True)
        if not args.clave_recursos.exists():
            with args.clave_recursos.open('xb') as f: f.write(secrets.token_bytes(32))
        clave=args.clave_recursos.read_bytes()
    if args.fuente=='nominal': out=nominal(args.archivo,args.corte,args.ambito,clave)
    elif args.fuente=='cubos':
        if not args.anio: parser.error('Cubos requiere --anio del libro.')
        out=cubos(args.archivo,args.corte,args.ambito,args.anio)
    else: out=sinba(args.archivo,args.ambito,clave)
    if not out['filas']: raise ValueError('La fuente no produjo filas; no publicar un corte vacío.')
    args.salida.mkdir(parents=True,exist_ok=True)
    name=f"{out['fuente']}_{args.ambito}_{out['corte']}_{out['sha256'][:16]}.json"
    dest=args.salida/name
    encoded=json.dumps(out,ensure_ascii=False,sort_keys=True,separators=(',',':'))
    if dest.exists() and dest.read_text(encoding='utf-8')!=encoded: raise ValueError('El archivo ya existe con otra transformación; conservar y revisar.')
    if not dest.exists(): dest.write_text(encoded,encoding='utf-8')
    print(json.dumps({'fuente':out['fuente'],'corte':out['corte'],'filas_agregadas':len(out['filas']),'auditoria':out['auditoria'],'salida':str(dest)},ensure_ascii=False))


if __name__=='__main__':
    try: main()
    except (ValueError,KeyError,StopIteration) as error: raise SystemExit(str(error))
